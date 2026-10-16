"""Đọc file Excel và chuẩn hóa thành các bảng (Block) để đưa cho LLM / code pandas.

Ý tưởng chính:
- Điền giá trị ô gộp (merged) vào mọi ô trong vùng gộp (trừ dòng tiêu đề/chú thích gộp ngang từ cột A).
- Chuyển chuỗi số ("1,234", "12.5") thành số; ký hiệu thiếu dữ liệu ("–", "�") thành None.
- Mỗi sheet được tách thành các khối: dòng tiêu đề (title) -> dòng header -> dòng dữ liệu -> chú thích.
- DataFrame dùng chỉ số = số dòng Excel, tên cột = chữ cái cột Excel; mô tả cột nằm ở `Block.columns`.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

MISSING_TOKENS = {"", "-", "–", "—", "…", "...", "�", "na", "n/a", "(na)", "(x)", "x", "(z)", "z", "nan", "none"}
NUM_RE = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
SHEET_PART = r"(?:'([^']+)'!|([A-Za-z0-9_\.]+)!)?"
RANGE_RE = re.compile(SHEET_PART + r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
REF_RE = re.compile(SHEET_PART + r"\$?([A-Z]{1,3})\$?(\d+)(?![\d(])")
FUNCS = {
    "SUM": lambda *a: sum(_flat(a)),
    "AVERAGE": lambda *a: sum(_flat(a)) / max(1, len(_flat(a))),
    "MIN": lambda *a: min(_flat(a)),
    "MAX": lambda *a: max(_flat(a)),
    "COUNT": lambda *a: len(_flat(a)),
    "ABS": abs,
    "ROUND": lambda x, n=0: round(x, int(n)),
}


def _flat(args) -> list:
    out = []
    for a in args:
        out.extend(_flat(a) if isinstance(a, (list, tuple)) else [a])
    return [x for x in out if is_num(x)]


HYPHEN_PREFIXES = {"self", "non", "part", "full", "year", "co", "ex", "multi", "semi", "well", "high", "low",
                   "long", "short", "post", "pre", "anti", "mid", "one", "two", "three", "four", "five"}


def _join_hyphen(m: re.Match) -> str:
    word = m.group(1)
    return f"{word}-" if word.lower() in HYPHEN_PREFIXES else word


def clean_text(s: str) -> str:
    s = s.replace("\r", "")
    s = re.sub(r"(\w+)-\n(?=[a-z])", _join_hyphen, s)  # "Con-\nstruction" -> "Construction", giữ "Self-employed"
    s = re.sub(r"\s*\n\s*", " ", s)
    return re.sub(r"[ \t]+", " ", s).strip()


def to_value(v):
    """Chuẩn hóa giá trị một ô."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        if isinstance(v, float):
            v = float(f"{v:.12g}")  # bỏ nhiễu dấu phẩy động: 1637.6799999999998 -> 1637.68
            return int(v) if v.is_integer() and abs(v) < 1e15 else v
        return v
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v
    s = clean_text(str(v))
    if s.lower() in MISSING_TOKENS:
        return None
    t = re.sub(r"(?<=\d),(?=\d{3}\b)", "", s).replace(" ", "")
    if NUM_RE.match(t):
        f = float(t)
        return int(f) if f.is_integer() and "." not in t and "e" not in t.lower() else f
    return s


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)


@dataclass
class Block:
    sheet: str
    index: int
    title: list[str]
    columns: dict[str, str]  # chữ cái cột -> mô tả header (đã làm phẳng)
    label_col: str
    df: pd.DataFrame  # index = số dòng Excel, columns = chữ cái cột
    raw: dict[int, list] = field(default_factory=dict)  # giá trị gốc theo dòng (để hiển thị)
    indents: dict[int, int] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    @property
    def name(self) -> str:
        return f"{self.sheet}#{self.index}"

    def render(self, max_rows: int | None = None) -> str:
        lines = [f"### Table {self.name} (sheet '{self.sheet}')"]
        if self.title:
            lines.append("Title: " + " / ".join(self.title))
        lines.append("Columns:")
        for c, desc in self.columns.items():
            tag = " (row label)" if c == self.label_col else ""
            lines.append(f"  {c}: {desc or '(no header)'}{tag}")
        cols = list(self.columns)
        lines.append("Rows (excel_row | " + " | ".join(cols) + "):")
        rows = list(self.raw.items())
        shown = rows if max_rows is None or len(rows) <= max_rows else rows[: max_rows - 10] + [None] + rows[-10:]
        for item in shown:
            if item is None:
                lines.append(f"... ({len(rows) - max_rows} rows omitted; full data is in the DataFrame) ...")
                continue
            r, row = item
            cells = [fmt_value(v) for v in row]
            ind = self.indents.get(r, 0)
            if ind:
                li = cols.index(self.label_col)
                cells[li] = "  " * ind + cells[li]
            while cells and cells[-1] == "":
                cells.pop()
            lines.append(f"{r} | " + " | ".join(cells))
        if self.notes:
            lines.append("Notes: " + " ".join(self.notes))
        return "\n".join(lines)


def _load_grid(ws_v, ws_f, wb_v):
    """Trả về (grid, indents, first_row, first_col) với ô gộp đã được điền."""
    min_r, min_c, max_r, max_c = ws_f.min_row, ws_f.min_column, ws_f.max_row, ws_f.max_column
    grid: list[list] = []
    indents: dict[int, int] = {}
    for r in range(min_r, max_r + 1):
        row = []
        for c in range(min_c, max_c + 1):
            v = ws_v.cell(r, c).value
            if v is None:
                f = ws_f.cell(r, c).value
                if isinstance(f, str) and f.startswith("="):
                    v = eval_formula(f, wb_v, ws_f.parent, ws_f.title)
            row.append(to_value(v))
            al = ws_f.cell(r, c).alignment
            if al is not None and al.indent and isinstance(row[-1], str):
                indents[r] = max(indents.get(r, 0), int(al.indent))
        grid.append(row)

    width = max_c - min_c + 1
    for mr in ws_f.merged_cells.ranges:
        r0, c0, r1, c1 = mr.min_row, mr.min_col, mr.max_row, mr.max_col
        val = grid[r0 - min_r][c0 - min_c]
        # Tiêu đề / chú thích gộp ngang từ cột đầu: chỉ giữ ở ô đầu.
        if r0 == r1 and c0 == min_c and (c1 - c0 + 1) >= max(2, int(0.6 * width)):
            continue
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if grid[r - min_r][c - min_c] is None:
                    grid[r - min_r][c - min_c] = val
    return grid, indents, min_r, min_c


def _cell(wb_v, wb_f, sheet: str, coord: str, depth: int):
    if sheet not in wb_v.sheetnames:
        raise KeyError(sheet)
    v = wb_v[sheet][coord].value
    if v is None:
        f = wb_f[sheet][coord].value
        if isinstance(f, str) and f.startswith("="):
            v = eval_formula(f, wb_v, wb_f, sheet, depth + 1)
    return to_value(v)


def eval_formula(formula: str, wb_v, wb_f, sheet: str, depth: int = 0):
    """Tính công thức đơn giản khi file không lưu sẵn giá trị: tham chiếu ô, + - * / ^, SUM/AVERAGE/MIN/MAX/COUNT."""
    if depth > 30:
        return None
    env: dict = {"__builtins__": {}, **FUNCS}

    def put(val) -> str:
        name = f"_v{len(env)}"
        env[name] = val
        return name

    def rng(m: re.Match) -> str:
        sh = m.group(1) or m.group(2) or sheet
        c1, r1, c2, r2 = m.group(3), int(m.group(4)), m.group(5), int(m.group(6))
        from openpyxl.utils import column_index_from_string as ci

        vals = [_cell(wb_v, wb_f, sh, f"{get_column_letter(c)}{r}", depth)
                for r in range(r1, r2 + 1) for c in range(ci(c1), ci(c2) + 1)]
        return put(vals)

    def ref(m: re.Match) -> str:
        v = _cell(wb_v, wb_f, m.group(1) or m.group(2) or sheet, f"{m.group(3)}{m.group(4)}", depth)
        return put(0 if v is None else v)

    try:
        expr = RANGE_RE.sub(rng, formula.lstrip("="))
        expr = REF_RE.sub(ref, expr).replace("^", "**")
        if re.search(r"[A-Za-z_]\w*\(", expr) and not all(f in FUNCS for f in re.findall(r"([A-Za-z_]\w*)\(", expr)):
            return None
        val = eval(expr, env)  # noqa: S307 - biểu thức đã được giới hạn
        return to_value(val) if is_num(val) else None
    except Exception:
        return None


def _year_only(vals) -> bool:
    nums = [v for v in vals if v is not None]
    return bool(nums) and all(isinstance(v, int) and 1900 <= v <= 2100 for v in nums)


def _split_blocks(sheet: str, grid, indents, r0: int, c0: int) -> list[Block]:
    ncols = max((len(r) for r in grid), default=0)
    letters = [get_column_letter(c0 + j) for j in range(ncols)]

    def nonempty(row):
        return [j for j, v in enumerate(row) if v is not None]

    def is_data(i):
        row = grid[i]
        rest = row[1:]
        return row[0] is not None and any(is_num(v) or isinstance(v, (dt.date, dt.datetime)) for v in rest) and not (
            i < 10 and _year_only(rest) and not isinstance(row[0], (int, float))
        )

    def header_like(k):  # >= 2 ô, toàn chữ (không có số/ngày)
        ne = nonempty(grid[k])
        return len(ne) >= 2 and all(isinstance(grid[k][j], str) for j in ne)

    def next_nonempty(k):
        k += 1
        while k < len(grid) and not nonempty(grid[k]):
            k += 1
        return k if k < len(grid) else None

    def starts_new_table(k):
        """Bảng mới xếp chồng: sau dòng trống là (tiêu đề một ô ->) dòng header toàn chữ, rồi tới dữ liệu."""
        if k == 0 or nonempty(grid[k - 1]):
            return False
        ne = nonempty(grid[k])
        if ne == [0]:
            nk = next_nonempty(k)
            return nk is not None and header_like(nk)
        if header_like(k):
            nk = next_nonempty(k)
            return nk is not None and is_data(nk)
        return False

    blocks: list[Block] = []
    i, n = 0, len(grid)
    pending_title: list[str] = []
    while i < n:
        # Dòng tiêu đề: chỉ có ô đầu tiên.
        header_rows: list[int] = []
        while i < n and not is_data(i):
            ne = nonempty(grid[i])
            if not ne:
                i += 1
                continue
            if ne == [0] and not header_rows:
                pending_title.append(str(grid[i][0]))
            else:
                header_rows.append(i)
            i += 1
        if i >= n:
            break
        data_rows: list[int] = []
        while header_rows and nonempty(grid[header_rows[-1]]) == [0] and len(header_rows) > 1:
            data_rows.insert(0, header_rows.pop())  # dòng nhóm (vd "TOTAL") ngay trước dữ liệu
        while i < n:
            row = grid[i]
            ne = nonempty(row)
            if ne and row[0] is None:  # header mới giữa vùng dữ liệu -> khối mới
                break
            if data_rows and starts_new_table(i):
                break
            data_rows.append(i)
            i += 1
        # Chú thích ở cuối: các dòng chỉ có cột đầu sau dòng số cuối cùng.
        last_num = max((k for k in data_rows if is_data(k)), default=None)
        notes = []
        if last_num is not None:
            tail = [k for k in data_rows if k > last_num]
            data_rows = [k for k in data_rows if k <= last_num]
            notes = [str(grid[k][0]) for k in tail if grid[k][0] is not None]
        if not data_rows:
            continue

        used = [j for j in range(ncols) if any(grid[k][j] is not None for k in data_rows)]
        columns: dict[str, str] = {}
        for j in used:
            parts: list[str] = []
            for h in header_rows:
                v = grid[h][j]
                if v is not None:
                    s = fmt_value(v)
                    if not parts or parts[-1] != s:
                        parts.append(s)
            columns[letters[j]] = " > ".join(parts)
        label_col = letters[used[0]]
        for j in used:  # cột nhãn = cột chữ đầu tiên
            if sum(isinstance(grid[k][j], str) for k in data_rows) > len(data_rows) / 2:
                label_col = letters[j]
                break
        raw = {r0 + k: [grid[k][j] for j in used] for k in data_rows if any(grid[k][j] is not None for j in used)}
        df = pd.DataFrame(list(raw.values()), index=list(raw), columns=[letters[j] for j in used], dtype=object)
        for c in df.columns:
            vals = df[c].dropna()
            if len(vals) and all(is_num(v) for v in vals):
                df[c] = pd.to_numeric(df[c])
            elif len(vals) and all(isinstance(v, (dt.date, dt.datetime)) for v in vals):
                df[c] = pd.to_datetime(df[c])
        df.index.name = "excel_row"
        blk_indents = {r: v for r, v in indents.items() if r in df.index}
        blocks.append(Block(sheet, len(blocks) + 1, pending_title, columns, label_col, df, raw, blk_indents, notes))
        pending_title = []
    return blocks


def load_workbook_blocks(path: str) -> list[Block]:
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    blocks: list[Block] = []
    for name in wb_f.sheetnames:
        ws_f, ws_v = wb_f[name], wb_v[name]
        if ws_f.max_row < 1 or ws_f.max_column < 1:
            continue
        grid, indents, r0, c0 = _load_grid(ws_v, ws_f, wb_v)
        blocks.extend(_split_blocks(name, grid, indents, r0, c0))
    return blocks


def render_blocks(blocks: list[Block], max_rows: int | None = None) -> str:
    return "\n\n".join(b.render(max_rows) for b in blocks)
