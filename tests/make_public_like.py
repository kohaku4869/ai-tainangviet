"""Sinh bộ test mô phỏng sát bộ public (p4_t01..p4_t10): xlsx + queries.json + answers.csv, đáp án tính bằng code.

    uv run python tests/make_public_like.py --out data/public_like
    uv run python main.py --data-dir data/public_like --gold data/public_like/answers.csv

Mỗi file bắt chước một bảng public (bảng hồ sơ người có header 2 tầng, bảng thương mại xếp chồng có công thức,
bảng CPI có cột Indent Level, các bảng BLS header gộp nhiều tầng, nhóm TOTAL/Men/Women/race, footnote,
ký hiệu '�'), nhưng số liệu được sinh ngẫu nhiên và nhất quán (tổng = tổng các thành phần, tỷ lệ tính lại từ số đếm).
Câu hỏi dùng đúng văn phong và yêu cầu định dạng đáp án của bộ public.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import random
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

R = random.Random(20261016)
Q: list[dict] = []
A: dict[str, str] = {}

INT = "Answer with only the integer, without commas or units."
DEC2 = "Answer with only the number rounded to two decimal places."
YN = "Answer with only Yes or No."
WRAP = Alignment(wrap_text=True, vertical="top")


def add(fname: str, query: str, answer) -> None:
    qid = f"pl_q{len(Q) + 1:03d}"
    Q.append({"id": qid, "file_name": fname, "query": query})
    A[qid] = str(answer)


def dec(x, n: int = 2) -> str:
    """Làm tròn half-up; từ chối trường hợp đúng biên .5 để đáp án không phụ thuộc kiểu làm tròn."""
    d = Decimal(str(x)) if not isinstance(x, Decimal) else x
    q = Decimal(1).scaleb(-n)
    if (d / q) % 1 == Decimal("0.5"):
        raise ValueError(f"tie when rounding {d} to {n} decimals; đổi seed")
    return str(d.quantize(q, rounding=ROUND_HALF_UP))


def mean(vals) -> Decimal:
    vals = [Decimal(str(v)) for v in vals]
    return sum(vals) / len(vals)


def argmax(d: dict):
    best = max(d.values())
    keys = [k for k, v in d.items() if v == best]
    assert len(keys) == 1, f"argmax không duy nhất: {keys}"
    return keys[0]


def r1(x: float) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def put(ws, r, c, v, r2=None, c2=None):
    ws.cell(r, c, v).alignment = WRAP
    if (r2 and r2 != r) or (c2 and c2 != c):
        ws.merge_cells(start_row=r, start_column=c, end_row=r2 or r, end_column=c2 or c)


def title(ws, text: str, ncol: int, unit: str = "[Numbers in thousands]"):
    put(ws, 1, 1, text, 1, ncol)
    ws.cell(2, 1, unit)


def footnotes(ws, lines: list[str], ncol: int):
    r = ws.max_row + 2
    for line in lines:
        put(ws, r, 1, line, r, ncol)
        r += 1


def row(ws, r, label, vals=None, indent=0):
    ws.cell(r, 1, label).alignment = Alignment(indent=indent)
    for j, v in enumerate(vals or []):
        ws.cell(r, 2 + j, v)


def split(total: int, weights: list[float]) -> list[int]:
    """Chia một số nguyên thành các phần theo trọng số, tổng đúng bằng total."""
    s = sum(weights)
    parts = [int(total * w / s) for w in weights]
    parts[R.randrange(len(parts))] += total - sum(parts)
    return parts


# ---------------------------------------------------------------- t01: hồ sơ người, header 2 tầng, ngày tháng
def t01(out: Path):
    fn = "pl_t01.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    for c, v in ((1, "No"), (2, "Name"), (5, "Date of"), (7, "Score")):
        ws.cell(1, c, v).font = Font(bold=True)
    for c, v in ((2, "First"), (3, "Middle"), (4, "Last"), (5, "Birth"), (6, "Joining")):
        ws.cell(2, c, v).font = Font(bold=True)
    for rng in ("A1:A2", "B1:D1", "E1:F1", "G1:G2"):
        ws.merge_cells(rng)
    firsts = ["An", "Binh", "Chi", "Dung", "Giang", "Ha", "Hieu", "Hoa", "Khanh", "Lan", "Linh", "Long", "Mai",
              "Minh", "Nam", "Ngoc", "Phuc", "Quang", "Son", "Thao", "Trang", "Tuan", "Viet", "Yen"]
    middles = ["Van", "Thi", "Minh", "Gia", "Hoang", "Ngoc", "Duc", "Thanh", "Quoc", "Bao"]
    lasts = ["Nguyen", "Tran", "Le", "Pham", "Hoang", "Vu", "Vo", "Dang", "Bui", "Do", "Ho", "Ngo", "Duong",
             "Ly", "Trinh", "Dinh", "Mai", "Cao", "Lam", "Phan", "Truong", "Ha", "Ta", "Luu"]
    R.shuffle(lasts)
    people = []
    for i in range(24):
        birth = dt.date(1982, 1, 1) + dt.timedelta(days=R.randint(0, 9000))
        join = dt.date(2021, 1, 1) + dt.timedelta(days=R.randint(0, 1450))
        people.append(dict(no=i + 1, first=R.choice(firsts), middle=R.choice(middles), last=lasts[i],
                           birth=birth, join=join, score=R.randint(45, 99)))
    for p in people:
        r = ws.max_row + 1
        ws.append([p["no"], p["first"], p["middle"], p["last"], p["birth"], p["join"], p["score"]])
        ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = "yyyy-mm-dd"
        ws.cell(r, 7).number_format = "0"
    age_at_join = {p["no"]: (p["join"] - p["birth"]).days for p in people}
    assert len(set(age_at_join.values())) == len(people)
    by_age = sorted(people, key=lambda p: age_at_join[p["no"]])  # trẻ -> già tại ngày gia nhập
    add(fn, "Who was the oldest person on their joining date? Answer with only the last name.", by_age[-1]["last"])
    add(fn, "What score did the third-youngest person (age measured on their joining date) receive? "
            "Answer with only the integer score.", by_age[2]["score"])
    by_birth = sorted(people, key=lambda p: p["birth"])  # già -> trẻ theo ngày sinh
    add(fn, "Among the five people with the earliest dates of birth, what was the lowest score? "
            "Answer with only the integer score.", min(p["score"] for p in by_birth[:5]))
    add(fn, f"How many people joined in 2023? {INT}", sum(p["join"].year == 2023 for p in people))
    born = [p["score"] for p in people if p["birth"] < dt.date(1995, 1, 1)]
    add(fn, f"What is the average score of the people born before 1995? {DEC2}", dec(mean(born)))
    wb.save(out / fn)


# ---------------------------------------------------------------- t02: 2 bảng xếp chồng, công thức, số thực
def t02(out: Path):
    fn = "pl_t02.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Monsumtable"
    years = [2020, 2021, 2022, 2023]
    data = {}
    for blk, top in (("Fiscal", 1), ("Calendar", 9)):
        put(ws, top, 1, None, top + 2, 1)
        put(ws, top, 2, f"{blk} years 2/" if blk == "Fiscal" else "Calendar years", top, 5)
        put(ws, top, 6, f"{blk} year-to-date", top, 7)
        ws.cell(top, 8, "August" if blk == "Fiscal" else "=H1")
        for j, y in enumerate(years + [2023, 2024, 2024]):
            ws.cell(top + 1, 2 + j, y)
        put(ws, top + 2, 2, "Billion dollars", top + 2, 8)
        exp = [R.uniform(135, 200) for _ in years] + [R.uniform(150, 175), R.uniform(150, 175), R.uniform(12, 18)]
        imp = [R.uniform(140, 205) for _ in years] + [R.uniform(150, 180), R.uniform(150, 180), R.uniform(12, 18)]
        exp = [round(v, 9) for v in exp]
        imp = [round(v, 9) for v in imp]
        data[blk] = (exp, imp)
        ws.cell(top + 3, 1, "Agricultural exports")
        ws.cell(top + 4, 1, "Agricultural imports")
        ws.cell(top + 6, 1, "Trade balance 3/")
        for j in range(7):
            col = get_column_letter(2 + j)
            ws.cell(top + 3, 2 + j, exp[j]).number_format = "0.000"
            ws.cell(top + 4, 2 + j, imp[j]).number_format = "0.000"
            # khối Calendar lưu cân đối thương mại dạng công thức (không có cached value)
            bal = f"={col}{top + 3}-{col}{top + 4}" if blk == "Calendar" else round(exp[j] - imp[j], 9)
            ws.cell(top + 6, 2 + j, bal).number_format = "0.000"
    fe, fi = data["Fiscal"]
    ce, ci = data["Calendar"]
    add(fn, "Are agricultural imports for the calendar year-to-date 2024 greater than those for the calendar "
            f"year-to-date 2023? {YN}", "Yes" if ci[5] > ci[4] else "No")
    bal = {y: fe[j] - fi[j] for j, y in enumerate(years)}
    add(fn, "Which fiscal year from 2020 to 2023 had the largest agricultural trade balance? "
            "Answer with only the year.", argmax(bal))
    add(fn, f"What was the agricultural trade balance for calendar year 2022 in billion dollars? {DEC2}",
        dec(round(ce[2] - ci[2], 9)))
    add(fn, "By how much did fiscal-year agricultural exports change from 2021 to 2023, in billion dollars? "
            f"Use a negative number for a decrease. {DEC2}", dec(round(fe[3] - fe[1], 9)))
    wb.save(out / fn)


# ---------------------------------------------------------------- t03: CPI, cột Indent Level + thụt lề
CPI_TREE = [
    (0, "All items"),
    (1, "Food"), (2, "Food at home"), (3, "Cereals and bakery products"), (3, "Meats, poultry, fish, and eggs"),
    (3, "Dairy and related products"), (3, "Fruits and vegetables"),
    (3, "Nonalcoholic beverages and beverage materials"), (3, "Other food at home"),
    (2, "Food away from home(1)"),
    None,
    (1, "Energy"), (2, "Energy commodities"), (3, "Fuel oil"), (3, "Motor fuel"), (4, "Gasoline (all types)"),
    (2, "Energy services"), (3, "Electricity"), (3, "Utility (piped) gas service"),
    None,
    (1, "All items less food and energy"), (2, "Commodities less food and energy commodities"),
    (3, "Apparel"), (3, "New vehicles"), (3, "Used cars and trucks"), (3, "Medical care commodities(1)"),
    (3, "Alcoholic beverages"), (3, "Tobacco and smoking products(1)"),
    (2, "Services less energy services"), (3, "Shelter"), (4, "Rent of primary residence"),
    (4, "Owners' equivalent rent of residences(2)"), (3, "Medical care services"), (4, "Physicians' services(1)"),
    (4, "Hospital services(1)(3)"), (3, "Transportation services"), (4, "Motor vehicle maintenance and repair(1)"),
    (4, "Motor vehicle insurance"), (4, "Airline fares"),
]


def clean(label: str) -> str:
    while label.endswith(")") and "(" in label and label[label.rindex("(") + 1:-1].isdigit():
        label = label[: label.rindex("(")]
    return label


def t03(out: Path):
    fn = "pl_t03.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    top = ["Indent Level", "Expenditure category", "Relative\nimportance\nAug.\n2024",
           "Unadjusted indexes", "Unadjusted indexes", "Unadjusted indexes",
           "Unadjusted percent change", "Unadjusted percent change",
           "Seasonally adjusted percent change", "Seasonally adjusted percent change",
           "Seasonally adjusted percent change"]
    sub = [None, None, None, "Aug.\n2023", "Jul.\n2024", "Aug.\n2024", "Aug.\n2023-\nAug.\n2024",
           "Jul.\n2024-\nAug.\n2024", "May\n2024-\nJun.\n2024", "Jun.\n2024-\nJul.\n2024", "Jul.\n2024-\nAug.\n2024"]
    for j, (a, b) in enumerate(zip(top, sub)):
        ws.cell(1, j + 1, a).alignment = WRAP
        ws.cell(2, j + 1, b).alignment = WRAP
    ws.cell(3, 1, "")
    ws.cell(3, 2, "")
    recs = []
    r = 4
    for node in CPI_TREE:
        if node is None:
            r += 1
            continue
        lvl, name = node
        ri = 100.0 if lvl == 0 else round(R.uniform(0.3, 40 / lvl), 3)
        if lvl and round(ri * 1000) % 10 == 5:
            ri = round(ri + 0.001, 3)
        i0 = round(R.uniform(130, 800), 3)
        i2 = round(i0 * R.uniform(0.9, 1.1), 3)
        i1 = round(i2 * R.uniform(0.97, 1.03), 3)
        yoy = r1((i2 / i0 - 1) * 100)
        mom = r1((i2 / i1 - 1) * 100)
        sa = [r1(R.uniform(-1.5, 1.5)) for _ in range(3)]
        vals = [lvl, name, ri, i0, i1, i2, yoy, mom, *sa]
        if "(3)" in name:  # dòng thiếu số liệu điều chỉnh mùa vụ như bộ public
            vals[8] = vals[9] = None
            sa = [None, None, sa[2]]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 1, v)
            if j >= 2 and v is not None:
                c.number_format = "#0.000" if j < 6 else "0.0"
        ws.cell(r, 2).alignment = Alignment(indent=lvl)
        recs.append(dict(lvl=lvl, name=clean(name), ri=ri, i0=i0, i1=i1, i2=i2, yoy=yoy, mom=mom, sa=sa))
        r += 1
    mom_abs = {x["name"]: abs(x["mom"]) for x in recs}
    add(fn, "Which category had the largest absolute unadjusted percent change from July 2024 to August 2024? "
            "Answer with only the category name.", argmax(mom_abs))
    shelter = next(x for x in recs if x["name"] == "Shelter")
    add(fn, f"What is the relative importance of Shelter in August 2024? {DEC2}", dec(shelter["ri"]))
    fah = recs[[x["name"] for x in recs].index("Food at home") + 1:][:6]
    add(fn, "Among the categories listed under 'Food at home', which had the highest unadjusted index in "
            "August 2024? Answer with only the category name.", argmax({x["name"]: x["i2"] for x in fah}))
    add(fn, "How many categories at indent level 4 had a negative unadjusted percent change from August 2023 "
            f"to August 2024? {INT}", sum(x["lvl"] == 4 and x["yoy"] < 0 for x in recs))
    diff = {x["name"]: abs(x["yoy"] - x["mom"]) for x in recs}
    add(fn, "Considering all categories, which category had the largest absolute difference between its "
            "12-month unadjusted percent change (Aug. 2023-Aug. 2024) and its 1-month unadjusted percent change "
            "(Jul. 2024-Aug. 2024)? Answer with only the category name.", argmax(diff))
    wb.save(out / fn)


# ---------------------------------------------------------------- t04: tình trạng việc làm × trình độ học vấn
EDU_COLS = ["Less than a high school diploma", "High school graduates, no college", "Some college or associate degree",
            "Some college, no degree", "Associate degree", "Bachelor's degree and higher", "Bachelor's degree only",
            "Advanced degree"]
STATUS = ["Civilian noninstitutional population", "Civilian labor force", "Participation rate", "Employed",
          "Employment-population ratio", "Unemployed", "Unemployment rate"]


def edu_block(scale: float):
    """Trả về dict status -> list 8 cột (6 lá + 2 cột tổng nhóm)."""
    pops, lfs, emps = [], [], []
    for base, part, ur in ((10000, .45, .06), (31000, .6, .04), (16000, .62, .033), (12000, .68, .028),
                           (27000, .72, .022), (17000, .74, .018)):
        p = int(base * scale * R.uniform(0.85, 1.15))
        lf = int(p * part * R.uniform(0.9, 1.1))
        e = lf - max(1, int(lf * ur * R.uniform(0.6, 1.6)))
        pops.append(p), lfs.append(lf), emps.append(e)

    def full(v):  # chèn cột tổng 'Some college or associate degree' và 'Bachelor's degree and higher'
        return [v[0], v[1], v[2] + v[3], v[2], v[3], v[4] + v[5], v[4], v[5]]

    pop, lf, emp = full(pops), full(lfs), full(emps)
    return block_from(pop, lf, emp)


def block_from(pop, lf, emp):
    un = [a - b for a, b in zip(lf, emp)]
    return {
        STATUS[0]: pop, STATUS[1]: lf, STATUS[2]: [r1(a / b * 100) for a, b in zip(lf, pop)],
        STATUS[3]: emp, STATUS[4]: [r1(a / b * 100) for a, b in zip(emp, pop)],
        STATUS[5]: un, STATUS[6]: [r1(a / b * 100) for a, b in zip(un, lf)],
    }


def t04(out: Path):
    fn = "pl_t04.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat07"
    title(ws, "HOUSEHOLD DATA\nANNUAL AVERAGES\n7. Employment status of the civilian noninstitutional population "
              "25 years and over by educational attainment, sex, race, and Hispanic or Latino ethnicity", 9)
    put(ws, 4, 1, "Employment status,\n sex, race, and Hispanic\n or Latino ethnicity", 6, 1)
    put(ws, 4, 2, 2024, 4, 9)
    put(ws, 5, 2, "Less than a\nhigh school\ndiploma", 6, 2)
    put(ws, 5, 3, "High school\ngraduates,\nno college(1)", 6, 3)
    put(ws, 5, 4, "Some college or associate degree", 5, 6)
    put(ws, 5, 7, "Bachelor's degree and higher", 5, 9)
    for c, v in ((4, "Total"), (5, "Some college,\nno degree"), (6, "Associate\ndegree"), (7, "Total(2)"),
                 (8, "Bachelor's\ndegree\nonly"), (9, "Advanced\ndegree")):
        ws.cell(6, c, v).alignment = WRAP
    men, women = edu_block(1.0), edu_block(1.0)
    total = block_from(*[[a + b for a, b in zip(men[s], women[s])] for s in (STATUS[0], STATUS[1], STATUS[3])])
    groups = {"TOTAL": total, "Men": men, "Women": women, "White": edu_block(0.75),
              "Black or African American": edu_block(0.12), "Asian": edu_block(0.07),
              "Hispanic or Latino ethnicity": edu_block(0.17)}
    r = 8
    for g, blk in groups.items():
        row(ws, r, g)
        r += 1
        for s in STATUS:
            row(ws, r, s, blk[s], indent=1)
            r += 1
    footnotes(ws, ["Footnotes:", "(1) Includes persons with a high school diploma or equivalent.",
                   "(2) Includes persons with bachelor's, master's, professional, and doctoral degrees.",
                   "NOTE: Estimates for the above race groups (White, Black or African American, and Asian) do not "
                   "sum to totals because data are not presented for all races."], 9)
    m = groups["Men"]
    add(fn, "What is the absolute difference between the civilian labor force and the number employed for men "
            f"with an 'Associate degree'? {INT}", abs(m[STATUS[1]][4] - m[STATUS[3]][4]))
    add(fn, f"What was the average labor force participation rate across all education levels for Asian? {DEC2}",
        dec(mean(groups["Asian"][STATUS[2]])))
    b = groups["Black or African American"][STATUS[6]]
    add(fn, "Which education level had the highest unemployment rate for Black or African American? "
            "Answer with only the education level.", argmax(dict(zip(EDU_COLS, b))))
    w = groups["Women"][STATUS[5]]
    add(fn, "For women, is the number of unemployed with an advanced degree greater than the number unemployed "
            f"with a bachelor's degree only? {YN}", "Yes" if w[7] > w[6] else "No")
    h = groups["Hispanic or Latino ethnicity"][STATUS[3]]
    add(fn, "What is the total number of employed Hispanic or Latino persons with less than a high school diploma "
            f"and high school graduates with no college combined? {INT}", h[0] + h[1])
    wb.save(out / fn)


# ---------------------------------------------------------------- t05: việc làm theo ngành × tuổi, nhóm race
INDUSTRIES = ["Mining, quarrying, and oil and gas extraction", "Construction", "Manufacturing",
              "Wholesale and retail trade", "Transportation and utilities", "Information", "Financial activities",
              "Professional and business services", "Education and health services", "Leisure and hospitality",
              "Other services(1)", "Public administration"]
IND_HDR = ["Mining, quarrying, and oil and gas extraction", "Con-\nstruction", "Manu-\nfacturing",
           "Wholesale and retail trade", "Trans-\nportation\nand utilities", "Information", "Financial\nactivities",
           "Professional and business services", "Education and health services", "Leisure and hospitality",
           "Other services(1)", "Public administration"]
IND_SIZE = [600, 12000, 15500, 19500, 9500, 2900, 11000, 20500, 36000, 14000, 7600, 7400]
AGE_ROWS = ["16 to 19 years", "20 years and over", "20 to 24 years", "25 years and over", "25 to 54 years",
            "55 years and over"]


def age_rows(leaf):  # leaf: dict age -> list per industry (16-19, 20-24, 25-54, 55+)
    def s(*ks):
        return [sum(leaf[k][i] for k in ks) for i in range(len(leaf["16-19"]))]
    return [s("16-19", "20-24", "25-54", "55+"), leaf["16-19"], s("20-24", "25-54", "55+"), leaf["20-24"],
            s("25-54", "55+"), leaf["25-54"], leaf["55+"]]


def t05(out: Path):
    fn = "pl_t05.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat14"
    title(ws, "HOUSEHOLD DATA\nANNUAL AVERAGES\n14. Employed persons in nonagricultural industries by age, sex, "
              "race, and Hispanic or Latino ethnicity", 13, "[In thousands]")
    put(ws, 4, 1, "Age, sex, race, and\n Hispanic or Latino\n ethnicity", 5, 1)
    put(ws, 4, 2, 2024, 4, 13)
    for j, h in enumerate(IND_HDR):
        ws.cell(5, 2 + j, h).alignment = WRAP

    def leaf(scale, female_share):
        out_ = {}
        for age, w in (("16-19", .035), ("20-24", .085), ("25-54", .63), ("55+", .25)):
            out_[age] = [max(0, int(sz * scale * female_share[i] * w * R.uniform(0.6, 1.4)))
                         for i, sz in enumerate(IND_SIZE)]
        return out_

    fem = [R.uniform(.1, .7) for _ in IND_SIZE]
    men_l, wom_l = leaf(1, [1 - f for f in fem]), leaf(1, fem)
    tot_l = {k: [a + b for a, b in zip(men_l[k], wom_l[k])] for k in men_l}
    sections = {"TOTAL": {"Total": tot_l, "Men": men_l, "Women": wom_l}}
    for race, sc in (("White", .77), ("Black or African American", .12), ("Asian", .065),
                     ("Hispanic or Latino ethnicity", .19)):
        sections[race] = {"Men": leaf(sc, [1 - f for f in fem]), "Women": leaf(sc, fem)}
    r = 7
    for sec, parts in sections.items():
        row(ws, r, sec)
        r += 1
        for who, lf in parts.items():
            rows = age_rows(lf)
            row(ws, r, f"{who}, 16 years and over", rows[0])
            for k, lab in enumerate(AGE_ROWS):
                row(ws, r + 1 + k, lab, rows[k + 1], indent=1)
            r += 7
        r += 1
    footnotes(ws, ["Footnotes:", "(1) Includes private households.",
                   "NOTE: Estimates for the above race groups (White, Black or African American, and Asian) do not "
                   "sum to totals because data are not presented for all races. Persons whose ethnicity is identified "
                   "as Hispanic or Latino may be of any race."], 13)
    names = [clean(n) for n in INDUSTRIES]
    w2554 = dict(zip(names, wom_l["25-54"]))
    add(fn, "How many more employed women aged 25 to 54 years were in Education and health services than in "
            f"Leisure and hospitality? {INT}", w2554["Education and health services"] - w2554["Leisure and hospitality"])
    add(fn, "Identify the industry with the highest number of employed men aged 55 years and over. "
            "Answer with only the industry name.", argmax(dict(zip(names, men_l["55+"]))))
    hw = dict(zip(names, sections["Hispanic or Latino ethnicity"]["Women"]["20-24"]))
    add(fn, "What is the total number of employed Hispanic or Latino women aged 20 to 24 years in the "
            f"Construction and Manufacturing industries combined? {INT}", hw["Construction"] + hw["Manufacturing"])
    am = age_rows(sections["Asian"]["Men"])[4]
    am = dict(zip(names, am))
    add(fn, "For Asian men aged 25 years and over, does employment in Information exceed that in Financial "
            f"activities? {YN}", "Yes" if am["Information"] > am["Financial activities"] else "No")
    tot = age_rows(tot_l)
    add(fn, "How many more employed people were aged 25 to 54 than aged 55 and over across all industries? "
            f"{INT}", sum(tot[5]) - sum(tot[6]))
    wb.save(out / fn)


# ---------------------------------------------------------------- t06: header gộp 4 tầng (cpsaat15)
AGES6 = ["16 to 17 years", "18 to 19 years", "20 to 24 years", "25 to 34 years", "35 to 44 years",
         "45 to 54 years", "55 to 64 years", "65 years and over"]


def t06(out: Path):
    fn = "pl_t06.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat15"
    put(ws, 1, 1, "Age and sex", 5, 1)
    put(ws, 1, 2, 2024, 1, 13)
    put(ws, 2, 2, "Agriculture and related industries", 2, 5)
    put(ws, 2, 6, "Nonagricultural industries", 2, 13)
    for c, v in ((2, "Total"), (3, "Wage and\nsalary\nworkers(1)"), (4, "Self-\nemployed\nworkers,\nunincor-\nporated"),
                 (5, "Unpaid\nfamily\nworkers"), (6, "Total"), (12, "Self-\nemployed\nworkers,\nunincor-\nporated"),
                 (13, "Unpaid\nfamily\nworkers")):
        put(ws, 3, c, v, 5, c)
    put(ws, 3, 7, "Wage and salary workers(1)", 3, 11)
    put(ws, 4, 7, "Total", 5, 7)
    put(ws, 4, 8, "Private industries", 4, 10)
    put(ws, 4, 11, "Govern-\nment", 5, 11)
    for c, v in ((8, "Total"), (9, "Private\nhousehold\nworkers"), (10, "Other\nprivate\nindustries")):
        ws.cell(5, c, v).alignment = WRAP
    ws.cell(6, 1, "")
    sizes = [.022, .03, .085, .22, .21, .2, .16, .07]

    def sex_leaf(scale):
        d = {}
        for a, s in zip(AGES6, sizes):
            ag = [int(1100 * scale * s * R.uniform(.5, 1.5) * k) for k in (.65, .3 + (.2 if "65" in a else 0), .02)]
            na = [int(150000 * scale * s * R.uniform(.8, 1.2) * k) for k in (.005, .78, .14, .07, .001)]
            d[a] = ag + na  # ag wage, ag self, ag unpaid, private household, other private, gov, self, unpaid
        return d

    men, wom = sex_leaf(1.0), sex_leaf(0.9)
    tot = {a: [x + y for x, y in zip(men[a], wom[a])] for a in AGES6}

    def cols(v):
        ag_w, ag_s, ag_u, ph, op, gov, self_, unp = v
        priv = ph + op
        wage = priv + gov
        return [ag_w + ag_s + ag_u, ag_w, ag_s, ag_u, wage + self_ + unp, wage, priv, ph, op, gov, self_, unp]

    def with_groups(d):
        out_ = {a: cols(v) for a, v in d.items()}
        s = lambda ks: [sum(out_[k][i] for k in ks) for i in range(12)]  # noqa: E731
        out_["16 to 19 years"] = s(AGES6[:2])
        out_["total"] = s(AGES6)
        return out_

    blocks = {"Total": with_groups(tot), "Men": with_groups(men), "Women": with_groups(wom)}
    r = 7
    for who, b in blocks.items():
        row(ws, r, f"{who}, 16 years and over", b["total"])
        r += 1
        for a in ["16 to 19 years"] + AGES6:
            row(ws, r, a, b[a], indent=2 if a in AGES6[:2] else 1)
            r += 1
        r += 1
    w = blocks["Women"]
    add(fn, "Calculate the total number of self-employed unincorporated workers in nonagricultural industries for "
            f"women aged 25 to 34 years and 35 to 44 years combined. {INT}", w["25 to 34 years"][10] + w["35 to 44 years"][10])
    m65 = blocks["Men"]["65 years and over"]
    add(fn, "What percentage of men aged 65 years and over working in agriculture and related industries were wage "
            f"and salary workers? {DEC2}", dec(Decimal(m65[1]) * 100 / Decimal(m65[0])))
    gov = {a: w[a][9] for a in ["16 to 19 years", *AGES6[2:]]}
    add(fn, "Which age group of women had the most government workers in nonagricultural industries? "
            "Answer with only the age group.", argmax(gov))
    add(fn, "Calculate the total number of unpaid family workers across all age groups in nonagricultural "
            f"industries. What is the final sum? {INT}", blocks["Total"]["total"][11])
    wb.save(out / fn)


# ---------------------------------------------------------------- t07: ngành × nghề, header 2 tầng
OCC = ["Management, business, and financial operations occupations", "Professional and related occupations",
       "Protective service occupations", "Service occupations, except protective",
       "Sales and related occupations", "Office and administrative support occupations",
       "Farming, fishing, and forestry occupations", "Construction and extraction occupations",
       "Installation, maintenance, and repair occupations", "Production occupations",
       "Transportation and material moving occupations"]
IND7 = [("Agriculture and related", 0), ("Mining, quarrying, and\noil and gas extraction", 0), ("Construction", 0),
        ("Manufacturing", 0), ("Durable goods", 1), ("Nondurable goods", 1), ("Wholesale and retail trade", 0),
        ("Wholesale trade", 1), ("Retail trade", 1), ("Transportation and utilities", 0), ("Information", 0),
        ("Financial activities", 0), ("Professional and business services", 0),
        ("Education and health services", 0), ("Leisure and hospitality", 0), ("Other services", 0),
        ("Public administration", 0)]


def t07(out: Path):
    fn = "pl_t07.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat17"
    title(ws, "HOUSEHOLD DATA\nANNUAL AVERAGES\n17. Employed persons by industry, sex, race, and occupation",
          13, "[In thousands]")
    put(ws, 4, 1, "Industry, sex, and race", 6, 1)
    put(ws, 4, 2, 2024, 4, 13)
    put(ws, 5, 2, "Total employed", 6, 2)
    for c1, c2, v in ((3, 4, "Management, professional, and related occupations"), (5, 6, "Service\noccupations"),
                      (7, 8, "Sales and office\noccupations"),
                      (9, 11, "Natural resources,\nconstruction, and\nmaintenance occupations"),
                      (12, 13, "Production, transportation, and material moving occupations")):
        put(ws, 5, c1, v, 5, c2)
    for j, o in enumerate(OCC):
        ws.cell(6, 3 + j, o).alignment = WRAP
    profile = {name: [R.uniform(.01, 1) for _ in OCC] for name, lvl in IND7}
    for name in ("Agriculture and related",):
        profile[name][6] *= 25
    for name in ("Construction",):
        profile[name][7] *= 20
    for name in ("Durable goods", "Nondurable goods"):
        profile[name][9] *= 12

    def sex_rows(scale):
        leaf = {}
        for name, lvl in IND7:
            if name in ("Manufacturing", "Wholesale and retail trade"):
                continue
            tot = int(R.uniform(500, 20000) * scale)
            leaf[name] = split(tot, profile[name])
        leaf["Manufacturing"] = [a + b for a, b in zip(leaf["Durable goods"], leaf["Nondurable goods"])]
        leaf["Wholesale and retail trade"] = [a + b for a, b in zip(leaf["Wholesale trade"], leaf["Retail trade"])]
        return leaf

    men, wom = sex_rows(1.0), sex_rows(0.8)
    tot = {k: [a + b for a, b in zip(men[k], wom[k])] for k in men}
    r = 8
    for sec, d in (("TOTAL", tot), ("Men", men), ("Women", wom)):
        row(ws, r, sec)
        r += 1
        for name, lvl in IND7:
            row(ws, r, name, [sum(d[name])] + d[name], indent=1 + lvl)
            r += 1
        r += 1
    footnotes(ws, ["NOTE: Estimates for the above race groups will not sum to totals shown because data are not "
                   "presented for all races."], 13)
    sel = ["Information", "Financial activities", "Public administration"]
    vals = {s: wom[s][5] for s in sel}
    assert len(set(vals.values())) == 3
    add(fn, "Rank the following sectors by the number of women employees in 'Office and administrative support "
            "occupations': Information, Financial activities, and Public administration. Answer with the three "
            "sector names in descending order, separated by a comma and one space, without numbering.",
        ", ".join(sorted(sel, key=vals.get, reverse=True)))
    add(fn, "Calculate the combined total of employees in 'Production occupations' and 'Transportation and "
            f"material moving occupations' for the Durable goods sector. {INT}",
        tot["Durable goods"][9] + tot["Durable goods"][10])
    add(fn, "Which occupation group has the largest number of employed men in the Construction industry? "
            "Answer with only the occupation name.", argmax(dict(zip(OCC, men["Construction"]))))
    add(fn, f"How many people in total are employed in Retail trade? {INT}", sum(tot["Retail trade"]))
    wb.save(out / fn)


# ---------------------------------------------------------------- t08: giờ làm việc + phân bố %, ký hiệu '�'
HOURS = ["1 to 4 hours", "5 to 14 hours", "15 to 29 hours", "30 to 34 hours", "35 to 39 hours", "40 hours",
         "41 to 48 hours", "49 to 59 hours", "60 hours and over"]


def t08(out: Path):
    fn = "pl_t08.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat19"
    put(ws, 1, 1, "Hours of work", 3, 1)
    put(ws, 1, 2, 2024, 1, 7)
    put(ws, 2, 2, "Persons at work", 2, 4)
    put(ws, 2, 5, "Percent distribution", 2, 7)
    for j, h in enumerate(["All\nindustries", "Agriculture\nand related\nindustries", "Nonagricultural\nindustries"] * 2):
        ws.cell(3, 2 + j, h).alignment = WRAP
    ws.cell(4, 1, "")
    ag = split(int(R.uniform(2000, 2400)), [R.uniform(.5, 1.5) * w for w in (2, 6, 12, 7, 4, 34, 7, 12, 17)])
    na = split(int(R.uniform(150000, 158000)), [R.uniform(.8, 1.2) * w for w in (1, 3.5, 11, 7, 6, 50, 8, 9, 5.5)])
    al = [a + b for a, b in zip(ag, na)]

    def grp(v):
        return {"Total, persons at work": sum(v), "1 to 34 hours": sum(v[:4]), **dict(zip(HOURS[:4], v[:4])),
                "35 hours and over": sum(v[4:]), **dict(zip(HOURS[4:6], v[4:6])), "41 hours and over": sum(v[6:]),
                **dict(zip(HOURS[6:], v[6:]))}

    g = [grp(al), grp(ag), grp(na)]
    pct = [{k: r1(v / d["Total, persons at work"] * 100) for k, v in d.items()} for d in g]
    r = 5
    for k in g[0]:
        row(ws, r, k, [d[k] for d in g] + [p[k] for p in pct],
            indent={"Total, persons at work": 0, "1 to 34 hours": 1, "35 hours and over": 1, "41 hours and over": 2}
            .get(k, 2 if k in HOURS[:6] else 3))
        r += 1
    avg = [(r1(R.uniform(38, 39)), r1(R.uniform(40, 43)), r1(R.uniform(38, 39))),
           (r1(R.uniform(41, 42.5)), r1(R.uniform(45, 48)), r1(R.uniform(41, 42.5)))]
    row(ws, r + 1, "Average hours, total at work", list(avg[0]) + ["�"] * 3)
    row(ws, r + 2, "Average hours, persons who usually work full time", list(avg[1]) + ["�"] * 3)
    add(fn, "What is the number of persons at work in nonagricultural industries who work 41 to 48 hours? "
            f"{INT}", g[2]["41 to 48 hours"])
    sel = ["41 to 48 hours", "49 to 59 hours", "60 hours and over"]
    add(fn, "Among the categories '41 to 48 hours', '49 to 59 hours', and '60 hours and over', which had the "
            "highest percentage of agricultural workers? Answer with only the hour category.",
        argmax({k: pct[1][k] for k in sel}))
    add(fn, "What are the average hours of persons who usually work full time in agriculture and related "
            "industries? Answer with only the number with one decimal place.", f"{avg[1][1]:.1f}")
    add(fn, "How many persons at work in all industries worked 40 hours or more? "
            f"{INT}", g[0]["40 hours"] + g[0]["41 hours and over"])
    wb.save(out / fn)


# ---------------------------------------------------------------- t09 / t10: giờ làm việc × lý do làm bán thời gian
def hours_header(ws, top: int, stub: str):
    put(ws, top, 1, stub, top + 3, 1)
    put(ws, top, 2, 2024, top, 9)
    put(ws, top + 1, 2, "Total\nat\nwork", top + 3, 2)
    put(ws, top + 1, 3, "Worked 1 to 34 hours", top + 1, 6)
    put(ws, top + 1, 7, "Worked\n35 hours\nor more", top + 3, 7)
    put(ws, top + 1, 8, "Average hours", top + 1, 9)
    put(ws, top + 2, 3, "Total", top + 3, 3)
    put(ws, top + 2, 4, "For\neconomic\nreasons", top + 3, 4)
    put(ws, top + 2, 5, "For noneconomic\nreasons", top + 2, 6)
    put(ws, top + 2, 8, "Total\nat\nwork", top + 3, 8)
    put(ws, top + 2, 9, "Persons who\nusually work\nfull time", top + 3, 9)
    ws.cell(top + 3, 5, "Usually work full time").alignment = WRAP
    ws.cell(top + 3, 6, "Usually work part time").alignment = WRAP


def hrs_leaf(total: int, pt_share: float):
    """[econ, noneconomic FT, noneconomic PT, 35+]"""
    pt = int(total * pt_share)
    return split(pt, [R.uniform(.08, .2), R.uniform(.15, .35), R.uniform(.4, .7)]) + [total - pt]


def hrs_row(v):
    return [sum(v), sum(v[:3]), *v]


def t09(out: Path):
    fn = "pl_t09.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat21"
    hours_header(ws, 1, "Industry and class of worker")
    ws.cell(5, 1, "")
    inds = ["Mining, quarrying, and oil and gas extraction", "Construction", "Durable goods", "Nondurable goods",
            "Wholesale and retail trade", "Transportation and utilities", "Information", "Financial activities",
            "Professional and business services", "Education and health services", "Leisure and hospitality",
            "Other services, except private households", "Private households", "Public administration"]
    sizes = [560, 10200, 9700, 5300, 18600, 9000, 2700, 10000, 18300, 33600, 13200, 5600, 650, 7800]
    pts = [.08, .14, .1, .11, .28, .17, .16, .14, .17, .24, .42, .26, .57, .15]
    leaf = {n: hrs_leaf(int(s * R.uniform(.9, 1.1)), p * R.uniform(.85, 1.15)) for n, s, p in zip(inds, sizes, pts)}
    avg = {n: (r1(R.uniform(34, 44) - p * 15), r1(R.uniform(40.5, 45.5))) for n, p in zip(inds, pts)}
    plus = lambda *ks: [sum(leaf[k][i] for k in ks) for i in range(4)]  # noqa: E731
    leaf["Manufacturing"] = plus("Durable goods", "Nondurable goods")
    leaf["Other services"] = plus("Other services, except private households", "Private households")
    for k in ("Manufacturing", "Other services"):
        avg[k] = (r1(R.uniform(37, 41)), r1(R.uniform(41, 44)))
    wage = plus(*inds)
    self_ = hrs_leaf(int(R.uniform(8000, 9000)), .38)
    unpaid = hrs_leaf(int(R.uniform(40, 70)), .55)
    total = [a + b + c for a, b, c in zip(wage, self_, unpaid)]
    order = ["Mining, quarrying, and oil and gas extraction", "Construction", "Manufacturing", "Durable goods",
             "Nondurable goods", "Wholesale and retail trade", "Transportation and utilities", "Information",
             "Financial activities", "Professional and business services", "Education and health services",
             "Leisure and hospitality", "Other services", "Other services, except private households",
             "Private households", "Public administration"]
    row(ws, 6, "Total, nonagricultural industries", hrs_row(total) + [38.6, 41.8])
    row(ws, 8, "Wage and salary workers(1)", hrs_row(wage) + [38.4, 41.6])
    r = 9
    for n in order:
        sub = n in ("Durable goods", "Nondurable goods", "Other services, except private households",
                    "Private households")
        row(ws, r, n, hrs_row(leaf[n]) + list(avg[n]), indent=2 if sub else 1)
        r += 1
    row(ws, r + 1, "Self-employed workers, unincorporated", hrs_row(self_) + [37.2, 46.0])
    row(ws, r + 3, "Unpaid family workers", hrs_row(unpaid) + [31.5, 44.9])
    footnotes(ws, ["(1) Includes self-employed workers whose businesses are incorporated."], 9)
    pb, eh = sum(leaf["Professional and business services"]), sum(leaf["Education and health services"])
    add(fn, "What is the absolute difference in the total number of wage and salary workers at work in "
            f"Professional and business services and Education and health services? {INT}", abs(pb - eh))
    add(fn, "How many wage and salary workers in Leisure and hospitality worked 1 to 34 hours for economic "
            f"reasons? {INT}", leaf["Leisure and hospitality"][0])
    add(fn, "Which industry had the highest average hours for persons who usually work full time? "
            "Answer with only the industry name.", argmax({n: avg[n][1] for n in order}))
    add(fn, "What is the number of self-employed unincorporated workers who usually work part time and worked "
            f"1 to 34 hours for noneconomic reasons? {INT}", self_[2])
    wb.save(out / fn)


def t10(out: Path):
    fn = "pl_t10.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cpsaat22"
    title(ws, "HOUSEHOLD DATA\nANNUAL AVERAGES \n22. Persons at work in nonagricultural industries by age, sex, "
              "marital status, and usual full- or part-time status", 9)
    hours_header(ws, 4, "Characteristic")
    ws.cell(4, 2, 2024)
    ages = [("16 to 17 years", 1000, .85), ("18 to 19 years", 1650, .62), ("20 to 24 years", 7500, .35),
            ("25 to 54 years", 52000, .17), ("55 years and over", 16500, .24)]
    sex = {}
    for s, sc, extra in (("Men", 1.0, 0.0), ("Women", .88, .1)):
        sex[s] = {a: hrs_leaf(int(n * sc * R.uniform(.9, 1.1)), min(.95, p + extra)) for a, n, p in ages}
    avg = {}

    def agg(s_list, a_list):
        return [sum(sex[s][a][i] for s in s_list for a in a_list) for i in range(4)]

    def ar(key):
        if key not in avg:
            avg[key] = (r1(R.uniform(20, 42)), r1(R.uniform(38, 44)))
        return list(avg[key])

    names = [a for a, _, _ in ages]
    lines = [("Total, 16 years and over", names, 0), ("16 to 19 years", names[:2], 1), ("16 to 17 years", names[:1], 2),
             ("18 to 19 years", names[1:2], 2), ("20 years and over", names[2:], 1), ("20 to 24 years", names[2:3], 2),
             ("25 years and over", names[3:], 2), ("25 to 54 years", names[3:4], 3),
             ("55 years and over", names[4:], 3)]
    r = 9
    row(ws, r, "AGE AND SEX")
    r += 1
    data = {}
    for prefix, sl in (("", ["Men", "Women"]), ("Men", ["Men"]), ("Women", ["Women"])):
        for lab, al, ind in lines:
            if prefix and lab in ("16 to 17 years", "18 to 19 years", "25 to 54 years", "55 years and over"):
                continue
            label = lab if not prefix else (f"{prefix}, 16 years and over" if lab.startswith("Total") else lab)
            v = agg(sl, al)
            data[(prefix or "All", lab)] = v
            row(ws, r, label, hrs_row(v) + ar((prefix, lab)), indent=ind)
            r += 1
        r += 1
    row(ws, r, "MARITAL STATUS")
    r += 1
    marital = {}
    for s in ("Men", "Women"):
        row(ws, r, s)
        r += 1
        tot = agg([s], names)
        parts = [split(x, [.55, .17, .28]) for x in tot]  # chia từng cột theo 3 tình trạng hôn nhân
        for k, lab in enumerate(["Married, spouse present(1)", "Widowed, divorced, or separated(2)", "Never married"]):
            v = [p[k] for p in parts]
            marital[(s, clean(lab))] = (v, ar((s, lab)))
            row(ws, r, lab, hrs_row(v) + ar((s, lab)), indent=1)
            r += 1
    footnotes(ws, ["Footnotes:", "(1) Refers to persons in both opposite-sex and same-sex married couples.",
                   "(2) Separated includes persons who are married, spouse absent.",
                   "NOTE: Full time is 35 hours or more per week; part time is less than 35 hours."], 9)
    v20 = data[("All", "20 years and over")]
    add(fn, "Based on the dataset, what are the numbers of people aged 20 years and over who worked 1 to 34 hours "
            "and who worked 35 hours or more? Answer as two integers in the order `1 to 34 hours, 35 hours or more`, "
            "separated by a comma and one space, without thousands separators or units.",
        f"{sum(v20[:3])}, {v20[3]}")
    nm, nw = marital[("Men", "Never married")][1][0], marital[("Women", "Never married")][1][0]
    add(fn, "What are the \"Average hours\" for total at work of never married men and women? Answer as "
            "`men, women`, using exactly two decimal places and separating the values with a comma and one space.",
        f"{nm:.2f}, {nw:.2f}")
    t = data[("All", "Total, 16 years and over")]
    add(fn, "What is the number of workers in \"Total, 16 years and over\" who usually work part time and worked "
            f"1 to 34 hours for noneconomic reasons? {INT}", t[2])
    wm = marital[("Women", "Married, spouse present")][0]
    add(fn, "How many married women with spouse present worked 1 to 34 hours for economic reasons? "
            f"{INT}", wm[0])
    wb.save(out / fn)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="data/public_like")
    ap.add_argument("--seed", type=int, default=20261016)
    args = ap.parse_args()
    out = Path(args.out)
    (out / "xlsx_files").mkdir(parents=True, exist_ok=True)
    # đáp án phải rõ ràng (không hòa khi argmax, không rơi đúng biên làm tròn) -> thử seed kế tiếp nếu vi phạm
    for seed in range(args.seed, args.seed + 200):
        R.seed(seed)
        Q.clear()
        A.clear()
        try:
            for g in (t01, t02, t03, t04, t05, t06, t07, t08, t09, t10):
                g(out / "xlsx_files")
            break
        except (AssertionError, ValueError):
            continue
    else:
        raise SystemExit("không tìm được seed hợp lệ")
    print(f"seed={seed}")
    (out / "queries.json").write_text(json.dumps(Q, indent=2, ensure_ascii=False), encoding="utf-8")
    with open(out / "answers.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "answer"])
        for k, v in A.items():
            w.writerow([k, v])
    print(f"{len(Q)} queries written to {out}")


if __name__ == "__main__":
    main()
