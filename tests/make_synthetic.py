"""Sinh bộ dữ liệu test tổng hợp (xlsx + queries.json + answers.csv) với đáp án tính bằng code.

    uv run python tests/make_synthetic.py --out data/synth

Phủ các dạng cấu trúc mà bộ public chưa có: nhiều sheet, nhiều bảng trên một sheet, bảng đặt cạnh nhau,
header gộp nhiều tầng, công thức, số dạng text có dấu phẩy, ký hiệu thiếu dữ liệu, phân cấp bằng thụt lề.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

R = random.Random(2026)
Q: list[dict] = []
A: dict[str, str] = {}

INT = "Answer with only the integer, without commas or units."
DEC2 = "Answer with only the number rounded to two decimal places."
YN = "Answer with only Yes or No."


def add(fname: str, query: str, answer) -> None:
    qid = f"s_q{len(Q) + 1:03d}"
    Q.append({"id": qid, "file_name": fname, "query": query})
    A[qid] = str(answer)


def bold(ws, *cells):
    for c in cells:
        ws[c].font = Font(bold=True)


# ---------------------------------------------------------------- s01: bảng giao dịch phẳng, có công thức
def s01(out: Path):
    fn = "s01_orders.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Order ID", "Order Date", "Region", "Product", "Quantity", "Unit Price", "Revenue"])
    regions = ["North", "South", "East", "West"]
    products = {"Laptop": 950, "Monitor": 210, "Keyboard": 35, "Mouse": 18, "Printer": 160}
    rows = []
    d0 = dt.date(2024, 1, 1)
    for i in range(1, 121):
        p = R.choice(list(products))
        price = round(products[p] * R.uniform(0.9, 1.1), 2)
        r = [f"ORD-{i:04d}", d0 + dt.timedelta(days=R.randint(0, 180)), R.choice(regions), p, R.randint(1, 50), price]
        rows.append(r)
        n = ws.max_row + 1
        ws.append(r + [f"=E{n}*F{n}"])  # công thức, không có cached value
    for r in rows:
        r.append(r[4] * r[5])
    last = ws.max_row
    ws.append([])
    ws.append(["Total", None, None, None, f"=SUM(E2:E{last})", None, f"=SUM(G2:G{last})"])

    rev = lambda f: sum(r[6] for r in rows if f(r))
    add(fn, f"What is the total revenue of the North region in March 2024? {DEC2}",
        f"{rev(lambda r: r[2] == 'North' and r[1].month == 3):.2f}")
    qty = {p: sum(r[4] for r in rows if r[3] == p) for p in products}
    add(fn, "Which product had the largest total quantity sold? Answer with only the product name.", max(qty, key=qty.get))
    add(fn, f"How many orders had a quantity greater than 40? {INT}", sum(r[4] > 40 for r in rows))
    lap = [r[5] for r in rows if r[3] == "Laptop"]
    add(fn, f"What is the average unit price of Laptop orders? {DEC2}", f"{sum(lap) / len(lap):.2f}")
    top = max(rows, key=lambda r: r[6])
    add(fn, "Which order ID has the highest revenue? Answer with only the order ID.", top[0])
    add(fn, f"What is the total quantity of all orders? {INT}", sum(r[4] for r in rows))
    wb.save(out / fn)


# ---------------------------------------------------------------- s02: nhiều sheet cùng bố cục
def s02(out: Path):
    fn = "s02_branches.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    branches = ["Ha Noi", "Hai Phong", "Da Nang", "Ho Chi Minh", "Can Tho"]
    data = {}
    for y in (2022, 2023, 2024):
        ws = wb.create_sheet(str(y))
        ws.append([f"Branch performance {y} (million VND)"])
        ws.append([])
        ws.append(["Branch", "Staff", "Revenue", "Cost", "Profit"])
        data[y] = {}
        for b in branches:
            rev, cost = R.randint(8000, 30000), R.randint(5000, 20000)
            staff = R.randint(20, 120)
            data[y][b] = (staff, rev, cost)
            n = ws.max_row + 1
            ws.append([b, staff, f"{rev:,}", f"{cost:,}", f"=C{n}-D{n}"])  # số dạng text có dấu phẩy
    notes = wb.create_sheet("Notes")
    notes.append(["Revenue and cost are in million VND."])
    notes.append(["Profit = Revenue - Cost."])

    g = data[2024]["Ha Noi"][1] - data[2022]["Ha Noi"][1]
    add(fn, f"By how much did the revenue of the Ha Noi branch change from 2022 to 2024? Answer with only the integer (negative if it decreased), without commas or units.", g)
    tc = {y: sum(v[2] for v in data[y].values()) for y in data}
    add(fn, "Which year had the highest total cost across all branches? Answer with only the year.", max(tc, key=tc.get))
    p23 = {b: v[1] - v[2] for b, v in data[2023].items()}
    add(fn, "Which branch had the highest profit in 2023? Answer with only the branch name.", max(p23, key=p23.get))
    st = sum(data[y]["Da Nang"][0] for y in data)
    add(fn, f"What is the total number of staff of the Da Nang branch summed over 2022, 2023 and 2024? {INT}", st)
    r24 = data[2024]["Can Tho"][1] / data[2024]["Can Tho"][0]
    add(fn, f"In 2024, what was the revenue per staff member of the Can Tho branch? {DEC2}", f"{r24:.2f}")
    add(fn, f"Was the 2023 revenue of Ho Chi Minh higher than its 2022 revenue? {YN}",
        "Yes" if data[2023]["Ho Chi Minh"][1] > data[2022]["Ho Chi Minh"][1] else "No")
    wb.save(out / fn)


# ---------------------------------------------------------------- s03: hai bảng xếp chồng trên một sheet
def s03(out: Path):
    fn = "s03_two_tables.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Warehouse"
    ws.append(["Table 1. Inventory by warehouse (units)"])
    ws.append(["Warehouse", "Rice", "Coffee", "Pepper", "Cashew"])
    inv = {}
    for w in ["WH-A", "WH-B", "WH-C", "WH-D"]:
        v = [R.randint(100, 900) for _ in range(4)]
        inv[w] = v
        ws.append([w] + v)
    ws.append([])
    ws.append([])
    ws.append(["Table 2. Shipments in Q1 2024 (units)"])
    ws.append(["Warehouse", "January", "February", "March"])
    shp = {}
    for w in ["WH-A", "WH-B", "WH-C", "WH-D"]:
        v = [R.randint(10, 300) for _ in range(3)]
        shp[w] = v
        ws.append([w] + v)
    add(fn, f"How many units of Coffee are stored in WH-C? {INT}", inv["WH-C"][1])
    add(fn, f"What is the total number of units shipped from WH-B in Q1 2024? {INT}", sum(shp["WH-B"]))
    feb = {w: v[1] for w, v in shp.items()}
    add(fn, "Which warehouse shipped the most units in February? Answer with only the warehouse name.", max(feb, key=feb.get))
    tot = {w: sum(v) for w, v in inv.items()}
    add(fn, "Which warehouse holds the largest total inventory across all products? Answer with only the warehouse name.",
        max(tot, key=tot.get))
    add(fn, f"What is the total inventory of Pepper across all warehouses? {INT}", sum(v[2] for v in inv.values()))
    wb.save(out / fn)


# ---------------------------------------------------------------- s04: header gộp 2 tầng, section, chú thích
def s04(out: Path):
    fn = "s04_trade_quarterly.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Trade"
    ws["A1"] = "Table 4. Trade value by commodity group and quarter"
    ws.merge_cells("A1:I1")
    ws["A2"] = "[Thousand USD]"
    ws["A4"] = "Commodity group"
    ws.merge_cells("A4:A5")
    col = 2
    quarters = ["Q1", "Q2", "Q3", "Q4"]
    for y in (2023, 2024):
        ws.cell(4, col, y)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 3)
        for i, q in enumerate(quarters):
            ws.cell(5, col + i, q)
        col += 4
    sections = {"Exports": ["Agricultural products", "Seafood", "Textiles", "Electronics"],
                "Imports": ["Machinery", "Fuel", "Chemicals", "Electronics"]}
    vals = {}
    r = 7
    for s, items in sections.items():
        ws.cell(r, 1, s)
        r += 1
        for it in items:
            ws.cell(r, 1, it)
            ws.cell(r, 1).alignment = Alignment(indent=1)
            v = []
            for j in range(8):
                x = R.randint(500, 9000) if not (it == "Fuel" and j == 3) else None
                v.append(x)
                ws.cell(r, 2 + j, f"{x:,}" if x is not None else "–")
            vals[(s, it)] = v
            r += 1
    r += 1
    ws.cell(r, 1, "Note: – data not available. Figures are preliminary.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)

    e = vals[("Exports", "Electronics")]
    i_ = vals[("Imports", "Electronics")]
    add(fn, f"What was the export value of Electronics in Q3 2024? {INT}", e[6])
    add(fn, f"What was the import value of Electronics in Q2 2023? {INT}", i_[1])
    sea = vals[("Exports", "Seafood")]
    add(fn, f"What is the total export value of Seafood in 2023? {INT}", sum(sea[:4]))
    q1 = {it: vals[("Exports", it)][4] for it in sections["Exports"]}
    add(fn, "Which export commodity group had the highest value in Q1 2024? Answer with only the commodity group name.",
        max(q1, key=q1.get))
    tex = vals[("Exports", "Textiles")]
    add(fn, f"What is the percentage change in Textiles exports from Q4 2023 to Q4 2024? {DEC2}",
        f"{(tex[7] - tex[3]) / tex[3] * 100:.2f}")
    fuel = [x for x in vals[("Imports", "Fuel")][:4] if x is not None]
    add(fn, f"What is the total import value of Fuel in 2023, ignoring quarters with no data? {INT}", sum(fuel))
    wb.save(out / fn)


# ---------------------------------------------------------------- s05: danh sách học viên, header gộp, tên tiếng Việt
def s05(out: Path):
    fn = "s05_students.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws["A1"], ws["B1"], ws["E1"], ws["G1"], ws["H1"] = "No", "Full name", "Date of", "Class", "Score"
    ws.merge_cells("A1:A2"), ws.merge_cells("B1:D1"), ws.merge_cells("E1:F1")
    ws.merge_cells("G1:G2"), ws.merge_cells("H1:H2")
    ws["B2"], ws["C2"], ws["D2"], ws["E2"], ws["F2"] = "Last", "Middle", "First", "Birth", "Enrollment"
    lasts = ["Nguyen", "Tran", "Le", "Pham", "Hoang", "Vu", "Dang", "Bui", "Do", "Ngo"]
    mids = ["Van", "Thi", "Minh", "Gia", "Ngoc", "Duc", "Thanh"]
    firsts = ["An", "Binh", "Chi", "Dung", "Giang", "Hai", "Khoa", "Lan", "Mai", "Nam", "Phuong", "Quan", "Son", "Trang"]
    people = []
    used = set()
    for i in range(1, 26):
        while True:
            b = dt.date(1995, 1, 1) + dt.timedelta(days=R.randint(0, 3650))
            e = dt.date(2019, 1, 1) + dt.timedelta(days=R.randint(0, 2000))
            if (e - b).days not in used and b not in [p[3] for p in people]:
                used.add((e - b).days)
                break
        p = [R.choice(lasts), R.choice(mids), R.choice(firsts), b, e, R.choice(["A1", "A2", "B1"]), R.randint(50, 100)]
        people.append(p)
        ws.append([i] + p)
        ws.cell(ws.max_row, 5).number_format = "dd/mm/yyyy"
        ws.cell(ws.max_row, 6).number_format = "dd/mm/yyyy"
    age = lambda p: (p[4] - p[3]).days
    s = sorted(people, key=age)
    add(fn, "Who was the youngest student on their enrollment date? Answer with only the first name.", s[0][2])
    add(fn, "Who was the third-oldest student on their enrollment date? Answer with only the last name.", s[-3][0])
    a2 = [p[6] for p in people if p[5] == "A2"]
    add(fn, f"What is the average score of students in class A2? {DEC2}", f"{sum(a2) / len(a2):.2f}")
    add(fn, f"How many students enrolled in 2021? {INT}", sum(p[4].year == 2021 for p in people))
    best = max(people, key=lambda p: (p[6], -p[3].toordinal()))
    hi = [p for p in people if p[6] == best[6]]
    if len(hi) == 1:
        add(fn, "What is the full name of the student with the highest score? Answer as 'Last Middle First'.",
            f"{best[0]} {best[1]} {best[2]}")
    born = sorted(people, key=lambda p: p[3])
    add(fn, "What score did the student with the earliest birth date receive? Answer with only the integer score.", born[0][6])
    wb.save(out / fn)


# ---------------------------------------------------------------- s06: hai bảng đặt cạnh nhau
def s06(out: Path):
    fn = "s06_side_by_side.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey"
    ws["A1"] = "Male respondents"
    ws["F1"] = "Female respondents"
    hdr = ["Age group", "Satisfied", "Neutral", "Unsatisfied"]
    for j, h in enumerate(hdr):
        ws.cell(2, 1 + j, h)
        ws.cell(2, 6 + j, h)
    groups = ["18-24", "25-34", "35-44", "45-54", "55+"]
    m, f = {}, {}
    for i, g in enumerate(groups):
        m[g] = [R.randint(20, 200) for _ in range(3)]
        f[g] = [R.randint(20, 200) for _ in range(3)]
        ws.cell(3 + i, 1, g)
        ws.cell(3 + i, 6, g)
        for j in range(3):
            ws.cell(3 + i, 2 + j, m[g][j])
            ws.cell(3 + i, 7 + j, f[g][j])
    add(fn, f"How many female respondents aged 25-34 were satisfied? {INT}", f["25-34"][0])
    add(fn, f"How many male respondents in total were unsatisfied? {INT}", sum(v[2] for v in m.values()))
    tot = {g: m[g][1] + f[g][1] for g in groups}
    add(fn, "Which age group had the most neutral respondents when combining males and females? Answer with only the age group.",
        max(tot, key=tot.get))
    wb.save(out / fn)


# ---------------------------------------------------------------- s07: phân cấp chỉ bằng thụt lề
def s07(out: Path):
    fn = "s07_budget.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Expense item", "Budget 2024", "Actual 2024"])
    tree = {"Personnel": ["Salaries", "Insurance", "Training"],
            "Operations": ["Rent", "Utilities", "Maintenance"],
            "Marketing": ["Online ads", "Events"]}
    leaf = {}
    for parent, kids in tree.items():
        kv = {k: (R.randint(50, 900), R.randint(50, 900)) for k in kids}
        leaf.update({(parent, k): v for k, v in kv.items()})
        ws.append([parent, sum(v[0] for v in kv.values()), sum(v[1] for v in kv.values())])
        for k, v in kv.items():
            ws.append([k, v[0], v[1]])
            ws.cell(ws.max_row, 1).alignment = Alignment(indent=1)
    ws.append(["Total", sum(v[0] for v in leaf.values()), sum(v[1] for v in leaf.values())])
    over = [k for (p, k), v in leaf.items() if v[1] > v[0]]
    add(fn, f"How many detailed expense items (not group totals) had actual spending above budget? {INT}", len(over))
    op = sum(v[1] for (p, k), v in leaf.items() if p == "Operations")
    add(fn, f"What was the actual 2024 spending on Operations? {INT}", op)
    diff = {k: abs(v[1] - v[0]) for (p, k), v in leaf.items()}
    add(fn, "Which detailed expense item had the largest absolute difference between budget and actual? Answer with only the item name.",
        max(diff, key=diff.get))
    tb = sum(v[0] for v in leaf.values())
    ta = sum(v[1] for v in leaf.values())
    add(fn, f"Was the total actual spending lower than the total budget? {YN}", "Yes" if ta < tb else "No")
    wb.save(out / fn)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="data/synth")
    args = ap.parse_args()
    out = Path(args.out)
    (out / "xlsx_files").mkdir(parents=True, exist_ok=True)
    for g in (s01, s02, s03, s04, s05, s06, s07):
        g(out / "xlsx_files")
    (out / "queries.json").write_text(json.dumps(Q, indent=2, ensure_ascii=False), encoding="utf-8")
    with open(out / "answers.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "answer"])
        for k, v in A.items():
            w.writerow([k, v])
    print(f"{len(Q)} queries written to {out}")


if __name__ == "__main__":
    main()
