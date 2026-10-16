"""Bộ công cụ khám phá workbook, được đưa vào namespace Python của LLM (CodeAct).

LLM gọi trực tiếp các hàm này trong code: overview(), view(), find(), grid(), table(), make_table(), final().
"""

from __future__ import annotations

import math
import re

import numpy as np
import pandas as pd

from ..utils.excel_reader import SheetData, WorkbookData
from ..utils.rendering import render_overview
from ..utils.table_detector import Table
from ..utils.values import fmt_value

_RANGE_RE = re.compile(r"^\$?([A-Z]{1,3})?\$?(\d+)?(?::\$?([A-Z]{1,3})?\$?(\d+)?)?$")

TOOL_DOCS = """Helper functions available in Python (sheet defaults to the first sheet):
- overview()                         -> print sheets, ranges and the guessed tables (column legend, warnings)
- view(sheet=None, cell_range=None, max_rows=40)
                                     -> print raw cells with Excel row numbers / column letters,
                                        e.g. view("Sheet1", "A1:H30"); merged cells are filled; indentation shown
- find(text, sheet=None, regex=False, limit=30)
                                     -> list of (sheet, cell, value) whose text contains `text` (case-insensitive)
- grid(sheet=None)                   -> DataFrame of the whole sheet: index = Excel row, columns = Excel letters
- table(name)                        -> DataFrame of a guessed table, e.g. table("Sheet1#1"):
                                        index = Excel row, column names = header texts exactly as in the overview
                                        (the column letter when a column has no header)
- make_table(sheet, header_rows, data_rows, columns=None)
                                     -> DataFrame built from YOUR reading of the layout, e.g.
                                        make_table("Sheet1", header_rows=[4, 5], data_rows="7:63", columns="A:I");
                                        column names = header texts joined with " > ", index = Excel row
- final(value)                       -> submit the final answer (number, string, or list for multi-part answers)
pandas is `pd`, numpy is `np`."""


def _parse_range(sd: SheetData, cell_range: str | None) -> tuple[int, int, int, int]:
    """'A1:H30' / 'A:C' / '5:20' -> (r1, c1, r2, c2) theo chỉ số Excel."""
    from openpyxl.utils import column_index_from_string as ci

    if not cell_range:
        return sd.min_row, sd.min_col, sd.max_row, sd.max_col
    m = _RANGE_RE.match(cell_range.replace(" ", "").upper())
    if not m:
        raise ValueError(f"Bad range {cell_range!r}; use e.g. 'A1:H30'")
    c1, r1, c2, r2 = m.groups()
    if m.group(0).count(":") == 0:
        c2, r2 = c1, r1
    return (int(r1) if r1 else sd.min_row, ci(c1) if c1 else sd.min_col,
            int(r2) if r2 else sd.max_row, ci(c2) if c2 else sd.max_col)


class WorkbookToolkit:
    def __init__(self, wb: WorkbookData, tables: list[Table]):
        self.wb = wb
        self.tables = tables
        self.final_called = False
        self.final_value = None

    # ------------------------------------------------------------------ helpers
    def _sheet(self, sheet: str | None) -> SheetData:
        if sheet is None:
            return next(iter(self.wb.sheets.values()))
        if sheet in self.wb.sheets:
            return self.wb.sheets[sheet]
        for name, sd in self.wb.sheets.items():  # khớp không phân biệt hoa thường
            if name.lower() == str(sheet).lower():
                return sd
        raise KeyError(f"No sheet {sheet!r}; sheets: {list(self.wb.sheets)}")

    # ------------------------------------------------------------------ tools
    def overview(self) -> None:
        print(render_overview(self.wb, self.tables))

    def view(self, sheet: str | None = None, cell_range: str | None = None, max_rows: int = 40) -> None:
        from openpyxl.utils import get_column_letter

        sd = self._sheet(sheet)
        r1, c1, r2, c2 = _parse_range(sd, cell_range)
        r1, c1 = max(r1, sd.min_row), max(c1, sd.min_col)
        r2, c2 = min(r2, sd.max_row), min(c2, sd.max_col)
        letters = [get_column_letter(c) for c in range(c1, c2 + 1)]
        print(f"Sheet '{sd.name}' {letters[0]}{r1}:{letters[-1]}{r2}")
        print("row | " + " | ".join(letters))
        shown = 0
        for r in range(r1, r2 + 1):
            vals = sd.grid[r - sd.min_row][c1 - sd.min_col: c2 - sd.min_col + 1]
            if all(v is None for v in vals):
                continue
            if shown >= max_rows:
                print(f"... (stopped after {max_rows} non-empty rows; call view with a later range)")
                break
            cells = []
            for j, v in enumerate(vals):
                ind = sd.indents.get((r, c1 + j), 0)
                cells.append("  " * ind + fmt_value(v))
            hidden = " (hidden)" if r in sd.hidden_rows else ""
            print(f"{r}{hidden} | " + " | ".join(cells))
            shown += 1
        merges = [m for m in sd.merges if _overlaps(m, r1, c1, r2, c2)]
        if merges:
            print("merged ranges: " + ", ".join(merges[:30]))

    def find(self, text: str, sheet: str | None = None, regex: bool = False, limit: int = 30) -> list:
        pat = re.compile(text if regex else re.escape(str(text)), re.I)
        sheets = [self._sheet(sheet)] if sheet else list(self.wb.sheets.values())
        hits = []
        for sd in sheets:
            for i, row in enumerate(sd.raw):
                for j, v in enumerate(row):
                    if v is not None and pat.search(fmt_value(v)):
                        hits.append((sd.name, f"{sd.letter(j)}{sd.min_row + i}", v))
                        if len(hits) >= limit:
                            return hits
        return hits

    def grid(self, sheet: str | None = None) -> pd.DataFrame:
        sd = self._sheet(sheet)
        df = pd.DataFrame(sd.grid, index=range(sd.min_row, sd.max_row + 1),
                          columns=[sd.letter(j) for j in range(sd.n_cols)], dtype=object)
        df = df.apply(lambda s: pd.to_numeric(s) if len(s.dropna()) and s.dropna().map(_is_number).all() else s)
        df.index.name = "excel_row"
        return df

    def _find_table(self, name: str) -> Table:
        key = str(name).strip().lower()
        for t in self.tables:
            if t.name.lower() == key:
                return t
        same_sheet = [t for t in self.tables if t.sheet.lower() == key]  # cho phép gọi bằng tên sheet
        if len(same_sheet) == 1:
            return same_sheet[0]
        raise KeyError(f"No table {name!r}; available: {[t.name for t in self.tables]}")

    def table(self, name: str) -> pd.DataFrame:
        """DataFrame của bảng đã đoán: index = số dòng Excel, tên cột = header (hoặc chữ cái cột nếu không có)."""
        t = self._find_table(name)
        names: list[str] = []
        for letter, desc in t.columns.items():
            col = desc or letter
            while col in names:
                col = f"{col} [{letter}]"
            names.append(col)
        df = t.df.copy()
        df.columns = names
        return df

    def make_table(self, sheet: str | None, header_rows, data_rows, columns: str | None = None) -> pd.DataFrame:
        sd = self._sheet(sheet)
        hdr = _rows(header_rows)
        r1, r2 = _row_span(data_rows)
        _, c1, _, c2 = _parse_range(sd, columns) if columns else (0, sd.min_col, 0, sd.max_col)
        cols = list(range(c1, c2 + 1))
        names: list[str] = []
        for c in cols:
            parts: list[str] = []
            for h in hdr:
                v = sd.grid[h - sd.min_row][c - sd.min_col] if sd.min_row <= h <= sd.max_row else None
                if v is not None and (not parts or parts[-1] != fmt_value(v)):
                    parts.append(fmt_value(v))
            name = " > ".join(parts) or sd.letter(c - sd.min_col)
            while name in names:
                name += "'"
            names.append(name)
        data = [[sd.grid[r - sd.min_row][c - sd.min_col] for c in cols] for r in range(r1, r2 + 1)
                if sd.min_row <= r <= sd.max_row]
        df = pd.DataFrame(data, index=[r for r in range(r1, r2 + 1) if sd.min_row <= r <= sd.max_row],
                          columns=names, dtype=object).dropna(how="all")
        df = df.apply(lambda s: pd.to_numeric(s) if len(s.dropna()) and s.dropna().map(_is_number).all() else s)
        df.index.name = "excel_row"
        return df

    def final(self, value) -> None:
        self.final_called = True
        self.final_value = value
        print(f"FINAL ANSWER SUBMITTED: {value!r}")

    def namespace(self) -> dict:
        """Namespace cho sandbox: công cụ + dữ liệu bảng đã đoán (bản sao)."""
        return {
            "pd": pd, "np": np, "math": math, "re": re,
            "overview": self.overview, "view": self.view, "find": self.find, "grid": self.grid,
            "table": self.table, "make_table": self.make_table, "final": self.final,
            "tables": {t.name: t.df.copy() for t in self.tables},
            "headers": {t.name: dict(t.columns) for t in self.tables},
            "df": self.tables[0].df.copy() if self.tables else None,
        }


def _is_number(v) -> bool:
    return isinstance(v, (int, float, np.number)) and not isinstance(v, bool)


def _rows(spec) -> list[int]:
    if isinstance(spec, int):
        return [spec]
    if isinstance(spec, str):
        a, _, b = spec.partition(":")
        return list(range(int(a), int(b or a) + 1))
    return [int(x) for x in spec]


def _row_span(spec) -> tuple[int, int]:
    rows = _rows(spec)
    return min(rows), max(rows)


def _overlaps(merge: str, r1: int, c1: int, r2: int, c2: int) -> bool:
    from openpyxl.utils.cell import range_boundaries

    mc1, mr1, mc2, mr2 = range_boundaries(merge)
    return not (mr2 < r1 or mr1 > r2 or mc2 < c1 or mc1 > c2)
