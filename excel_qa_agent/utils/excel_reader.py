"""Đọc workbook thành cấu trúc lưới thuần Python (giá trị đã chuẩn hóa + thông tin định dạng)."""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache

import openpyxl
from openpyxl.utils import get_column_letter

from .formulas import eval_formula
from .values import to_value


@dataclass
class SheetData:
    name: str
    min_row: int
    min_col: int
    grid: list[list]  # giá trị, ô gộp đã điền (trừ tiêu đề/chú thích gộp ngang từ cột đầu)
    raw: list[list]  # giá trị, ô gộp KHÔNG điền
    merges: list[str] = field(default_factory=list)
    indents: dict[tuple[int, int], int] = field(default_factory=dict)  # (row, col) -> mức thụt lề
    bold: set[tuple[int, int]] = field(default_factory=set)
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[str] = field(default_factory=set)
    formulas: dict[tuple[int, int], str] = field(default_factory=dict)

    @property
    def n_rows(self) -> int:
        return len(self.grid)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.grid), default=0)

    @property
    def max_row(self) -> int:
        return self.min_row + self.n_rows - 1

    @property
    def max_col(self) -> int:
        return self.min_col + self.n_cols - 1

    def letter(self, j: int) -> str:
        """Chữ cái cột cho chỉ số cột j (0-based trong lưới)."""
        return get_column_letter(self.min_col + j)

    @property
    def dimensions(self) -> str:
        return f"{get_column_letter(self.min_col)}{self.min_row}:{get_column_letter(self.max_col)}{self.max_row}"


@dataclass
class WorkbookData:
    path: str
    sheets: dict[str, SheetData]


def _read_sheet(ws_f, ws_v, wb_v, wb_f) -> SheetData:
    min_r, min_c, max_r, max_c = ws_f.min_row, ws_f.min_column, ws_f.max_row, ws_f.max_column
    raw: list[list] = []
    sd = SheetData(ws_f.title, min_r, min_c, [], [])
    for r in range(min_r, max_r + 1):
        row = []
        for c in range(min_c, max_c + 1):
            cell_f = ws_f.cell(r, c)
            v = ws_v.cell(r, c).value
            f = cell_f.value
            if isinstance(f, str) and f.startswith("="):
                sd.formulas[(r, c)] = f
                if v is None:
                    v = eval_formula(f, wb_v, wb_f, ws_f.title)
            row.append(to_value(v))
            if row[-1] is not None:
                al, font = cell_f.alignment, cell_f.font
                if al is not None and al.indent:
                    sd.indents[(r, c)] = int(al.indent)
                if font is not None and font.bold:
                    sd.bold.add((r, c))
        raw.append(row)

    grid = [list(r) for r in raw]
    width = max_c - min_c + 1
    for mr in ws_f.merged_cells.ranges:
        sd.merges.append(str(mr))
        r0, c0, r1, c1 = mr.min_row, mr.min_col, mr.max_row, mr.max_col
        val = raw[r0 - min_r][c0 - min_c]
        # Tiêu đề/chú thích gộp ngang từ cột đầu và trải gần hết chiều ngang: chỉ giữ ở ô đầu.
        if r0 == r1 and c0 == min_c and (c1 - c0 + 1) >= max(2, int(0.6 * width)):
            continue
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if grid[r - min_r][c - min_c] is None:
                    grid[r - min_r][c - min_c] = val

    sd.grid, sd.raw = grid, raw
    sd.hidden_rows = {r for r, d in ws_f.row_dimensions.items() if d.hidden}
    sd.hidden_cols = {k for k, d in ws_f.column_dimensions.items() if d.hidden}
    return sd


@lru_cache(maxsize=32)
def read_workbook(path: str) -> WorkbookData:
    """Đọc workbook (có cache theo đường dẫn vì nhiều câu hỏi dùng chung một file)."""
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    sheets = {}
    for name in wb_f.sheetnames:
        ws_f = wb_f[name]
        if ws_f.max_row < 1 or ws_f.max_column < 1:
            continue
        sheets[name] = _read_sheet(ws_f, wb_v[name], wb_v, wb_f)
    return WorkbookData(path, sheets)
